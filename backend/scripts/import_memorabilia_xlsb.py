"""
Idempotent import of Presidents' Memorabilia Collection into the artifacts catalog.

Supports:
  - .xlsb via Excel COM on Windows
  - .xlsx via openpyxl
  - .csv (columns: id_code,description,remarks,owner_name)

Usage (from backend/):
  python scripts/import_memorabilia_xlsb.py
  python scripts/import_memorabilia_xlsb.py --path ../data/imports/presidents_memorabilia_collection.xlsb
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from sqlalchemy import select, text  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from app.models import Artifact, President, User  # noqa: E402
from app.services.artifact_tags import preload_tag_cache, sync_auto_tags  # noqa: E402
from app.services.catalog_seed import _curator_account, _ensure_presidents  # noqa: E402
from app.services.facet_hints import apply_facet_hints_if_empty  # noqa: E402
from app.services.id_codes import (  # noqa: E402
    is_gallery_owner_label,
    looks_like_id_code,
    normalize_id_code,
    owner_for_id_code,
)
from app.services.photo_uploads import PHOTOS_ROOT, ensure_photo_dir  # noqa: E402

DEFAULT_PATH = Path(__file__).resolve().parents[2] / "data" / "imports" / "presidents_memorabilia_collection.xlsb"
PLACEHOLDER_IMAGE = "https://sulyap-lasalyano-gecapistranos-projects.vercel.app/splash/splash-logo.png"

SHEET_OWNER = {
    "andrew": "Br. Andrew B. Gonzalez, FSC",
    "gonzal": "Br. Andrew B. Gonzalez, FSC",
    "rafael": "Br. Rafael S. Donato, FSC",
    "donato": "Br. Rafael S. Donato, FSC",
    "rolando": "Br. Rolando V. Dizon, FSC",
    "dizon": "Br. Rolando V. Dizon, FSC",
    "augustine": "Br. Gus Boquer, FSC",
    "boquer": "Br. Gus Boquer, FSC",
    "francisco": 'Br. Francisco "Sockie" dela Rosa VI, FSC',
    "dela rosa": 'Br. Francisco "Sockie" dela Rosa VI, FSC',
    "riola": "Br. Iñigo Riola, FSC",
    "iñigo": "Br. Iñigo Riola, FSC",
    "inigo": "Br. Iñigo Riola, FSC",
}


def _owner_for_sheet(name: str) -> str:
    key = (name or "").lower()
    if is_gallery_owner_label(name) or "gallery" in key:
        return ""
    for needle, owner in SHEET_OWNER.items():
        if needle in key:
            return owner
    return ""


def _resolve_owner(*, id_code: str, sheet_owner: str = "", csv_owner: str = "") -> str:
    """Prefer accession prefix (RD→Donato, RDi→Dizon, …); never invent a gallery owner."""
    by_code = owner_for_id_code(id_code)
    if by_code:
        return by_code
    for candidate in (csv_owner, sheet_owner):
        label = (candidate or "").strip()
        if label and not is_gallery_owner_label(label):
            return label
    return ""


def title_from_description(description: str, *, max_len: int = 200) -> str:
    """Source workbook has no separate title column — derive a short title from description."""
    text = re.sub(r"\s+", " ", (description or "").strip())
    if not text:
        return ""
    # Prefer first clause; avoid splitting decimals like 6.5
    first = text
    for sep in ("; ", " — ", " - "):
        if sep in text:
            first = text.split(sep, 1)[0].strip()
            break
    else:
        m = re.search(r"(?<=\D)\.\s+", text)
        if m:
            first = text[: m.start() + 1].strip()
    if len(first) > max_len:
        cut = first[: max_len - 1].rsplit(" ", 1)[0].rstrip(",;:-")
        return (cut or first[: max_len - 1]).rstrip() + "…"
    return first.rstrip(",;:-")


def _export_shape_png(shape: object, dest: Path) -> bool:
    """xlsb Picture.Export is unreliable; Copy + clipboard works on Windows."""
    import time

    from PIL import ImageGrab  # type: ignore

    try:
        shape.Copy()
        time.sleep(0.05)
        img = ImageGrab.grabclipboard()
        if img is None:
            return False
        dest.parent.mkdir(parents=True, exist_ok=True)
        img.convert("RGB").save(dest, format="PNG")
        return dest.is_file() and dest.stat().st_size > 200
    except Exception:  # noqa: BLE001
        return False


def _read_xlsb_com(path: Path, *, export_photos: bool = False) -> list[dict]:
    import win32com.client  # type: ignore

    excel = win32com.client.Dispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    rows: list[dict] = []
    try:
        wb = excel.Workbooks.Open(str(path), False, True)
        if export_photos:
            ensure_photo_dir()
        for ws in wb.Worksheets:
            owner = _owner_for_sheet(ws.Name)
            used = ws.UsedRange
            if not used:
                continue
            n_rows = int(used.Rows.Count)
            shape_by_row: dict[int, object] = {}
            if export_photos:
                try:
                    for shape in ws.Shapes:
                        try:
                            top_row = int(shape.TopLeftCell.Row)
                            # Prefer leftmost / first picture on the row
                            if top_row not in shape_by_row:
                                shape_by_row[top_row] = shape
                        except Exception:  # noqa: BLE001
                            continue
                except Exception:  # noqa: BLE001
                    shape_by_row = {}
            for r in range(1, min(n_rows, 400) + 1):
                acc = str(ws.Cells(r, 1).Text or "").strip()
                if not looks_like_id_code(acc):
                    continue
                desc = str(ws.Cells(r, 2).Text or "").strip()
                remarks = str(ws.Cells(r, 4).Text or "").strip()
                image_url = PLACEHOLDER_IMAGE
                code = normalize_id_code(acc)
                if export_photos:
                    shape = shape_by_row.get(r) or shape_by_row.get(r - 1) or shape_by_row.get(r + 1)
                    if shape is not None:
                        dest = PHOTOS_ROOT / f"import_{code}.png"
                        if _export_shape_png(shape, dest):
                            image_url = f"/curator-media/artifacts/{dest.name}"
                rows.append(
                    {
                        "id_code": code,
                        "description": desc,
                        "remarks": remarks,
                        "owner_name": _resolve_owner(id_code=code, sheet_owner=owner),
                        "image_url": image_url,
                    }
                )
        wb.Close(False)
    finally:
        excel.Quit()
    return rows


def _read_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict] = []
    for ws in wb.worksheets:
        owner = _owner_for_sheet(ws.title)
        for row in ws.iter_rows(min_row=1, max_col=4, values_only=True):
            acc = str(row[0] or "").strip()
            if not looks_like_id_code(acc):
                continue
            code = normalize_id_code(acc)
            rows.append(
                {
                    "id_code": code,
                    "description": str(row[1] or "").strip(),
                    "remarks": str(row[3] or "").strip() if len(row) > 3 else "",
                    "owner_name": _resolve_owner(id_code=code, sheet_owner=owner),
                    "image_url": PLACEHOLDER_IMAGE,
                }
            )
    return rows


def _read_csv(path: Path) -> list[dict]:
    rows: list[dict] = []
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            acc = row.get("id_code") or row.get("accession") or row.get("ACCESSION NO.") or ""
            if not looks_like_id_code(acc):
                continue
            code = normalize_id_code(acc)
            rows.append(
                {
                    "id_code": code,
                    "description": (row.get("description") or row.get("DESCRIPTION") or "").strip(),
                    "remarks": (row.get("remarks") or row.get("REMARKS") or "").strip(),
                    "owner_name": _resolve_owner(
                        id_code=code,
                        csv_owner=(row.get("owner_name") or "").strip(),
                    ),
                    "image_url": PLACEHOLDER_IMAGE,
                }
            )
    return rows


def load_rows(path: Path, *, export_photos: bool = False) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".xlsb":
        return _read_xlsb_com(path, export_photos=export_photos)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(path)
    if suffix == ".csv":
        return _read_csv(path)
    raise SystemExit(f"Unsupported file type: {suffix}")


def upsert_rows(rows: list[dict], *, force_text: bool = False) -> tuple[int, int]:
    inserted = updated = 0
    with SessionLocal() as db:
        db.execute(text("SET statement_timeout = '120s'"))
        curator = _curator_account(db)
        _ensure_presidents(db)
        presidents = {p.name: p.id for p in db.scalars(select(President)).all()}
        tag_cache = preload_tag_cache(db)
        for i, row in enumerate(rows, start=1):
            code = row["id_code"]
            existing = db.scalar(select(Artifact).where(Artifact.id_code == code))
            owner_name = _resolve_owner(id_code=code, csv_owner=row.get("owner_name") or "")
            president_id = presidents.get(owner_name) if owner_name else None
            derived_title = title_from_description(row["description"])
            if existing is None:
                art = Artifact(
                    owner_id=curator.id,
                    title=derived_title,
                    description=row["description"],
                    image_url=row.get("image_url") or PLACEHOLDER_IMAGE,
                    id_code=code,
                    owner_name=owner_name,
                    year="",
                    year_acquired="",
                    subtitle="Presidents' Memorabilia Collection",
                    icon="🏺",
                    remarks=row.get("remarks") or "",
                    president_id=president_id,
                    published=True,
                )
                db.add(art)
                db.flush()
                apply_facet_hints_if_empty(art)
                sync_auto_tags(db, art, cache=tag_cache)
                inserted += 1
            else:
                if force_text or not (existing.description or "").strip():
                    existing.description = row["description"]
                if force_text or not (existing.title or "").strip():
                    existing.title = title_from_description(existing.description or row["description"])
                if force_text or not (existing.remarks or "").strip():
                    existing.remarks = row.get("remarks") or ""
                current_owner = (existing.owner_name or "").strip()
                should_fix_owner = force_text or not current_owner or is_gallery_owner_label(current_owner)
                if not should_fix_owner and owner_name and current_owner != owner_name:
                    # Accession prefix is authoritative (RD→Donato, RDi→Dizon).
                    by_code = owner_for_id_code(code)
                    if by_code and current_owner != by_code:
                        should_fix_owner = True
                if should_fix_owner and owner_name:
                    existing.owner_name = owner_name
                    existing.president_id = president_id
                if row.get("image_url") and (
                    force_text
                    or not existing.image_url
                    or existing.image_url == PLACEHOLDER_IMAGE
                    or existing.image_url.startswith("/curator-media/")
                ):
                    if row["image_url"] != PLACEHOLDER_IMAGE or existing.image_url == PLACEHOLDER_IMAGE:
                        existing.image_url = row["image_url"]
                apply_facet_hints_if_empty(existing)
                sync_auto_tags(db, existing, cache=tag_cache)
                updated += 1
            if i % 25 == 0:
                db.commit()
                print(f"  committed {i}/{len(rows)}")
        db.commit()
    return inserted, updated


def main() -> None:
    parser = argparse.ArgumentParser(description="Import memorabilia workbook into catalog DB")
    parser.add_argument("--path", type=Path, default=DEFAULT_PATH)
    parser.add_argument("--force-text", action="store_true")
    parser.add_argument("--export-photos", action="store_true", help="Export embedded Excel shapes (slow)")
    args = parser.parse_args()
    path = args.path.expanduser().resolve()
    if not path.is_file():
        raise SystemExit(f"File not found: {path}")
    print(f"Loading {path} ...")
    rows = load_rows(path, export_photos=args.export_photos)
    print(f"Parsed {len(rows)} accession rows")
    inserted, updated = upsert_rows(rows, force_text=args.force_text)
    print(f"Done. inserted={inserted} updated={updated}")


if __name__ == "__main__":
    main()
